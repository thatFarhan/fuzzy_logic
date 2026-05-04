import openpyxl

def read_file(path_file):
    wb = openpyxl.load_workbook(path_file)
    worksheet = wb.active

    data = []
    header = [cell.value for cell in worksheet[1]]

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, row))

        if record.get('id Pelanggan') is not None:
            data.append({
                'id'       : int(record['id Pelanggan']),
                'pelayanan': float(record['Pelayanan']),
                'harga'    : float(record['harga'])
            })
    return data

def linear_ascent(x, a, b):
    if x <= a:
        return 0.0
    elif a < x <= b:
        return (x - a) / (b - a)
    else:
        return 1.0

def linear_descent(x, a, b):
    if x >= b:
        return 0.0
    elif a <= x < b:
        return (b - x) / (b - a)
    else:
        return 1.0

def triangle(x, a, b, c):
    if x <= a or x >= c:
        return 0.0
    elif a < x <= b:
        return (x - a) / (b - a)
    elif b <= x < c:
        return (c - x) / (c - b)

def fuzzify_service(p):
    return {
        'buruk' : linear_descent(p, 1, 40),
        'sedang': triangle(p, 20, 50, 80),
        'baik'  : linear_ascent(p, 60, 100)
    }

def fuzzify_price(h):
    return {
        'murah' : linear_descent(h, 20000, 33000),
        'sedang': triangle(h, 26500, 37500, 50500),
        'mahal' : linear_ascent(h, 44000, 55000)
    }

def inference(mu_service, mu_price):
    service = mu_service
    price = mu_price

    r1 = min(service['baik'],   price['murah'])   # Sangat_Layak
    r2 = min(service['baik'],   price['sedang'])  # Layak
    r3 = min(service['baik'],   price['mahal'])   # Cukup_Layak
    r4 = min(service['sedang'], price['murah'])   # Layak
    r5 = min(service['sedang'], price['sedang'])  # Cukup_Layak
    r6 = min(service['sedang'], price['mahal'])   # Tidak_Layak
    r7 = min(service['buruk'],  price['murah'])   # Cukup_Layak
    r8 = min(service['buruk'],  price['sedang'])  # Tidak_Layak
    r9 = min(service['buruk'],  price['mahal'])   # Tidak_Layak

    return {
        'tidak_layak' : max(r6, r8, r9),
        'cukup_layak' : max(r3, r5, r7),
        'layak'       : max(r2, r4),
        'sangat_layak': r1
    }

def mu_output_tidak_layak(z):
    return linear_descent(z, 0, 35)

def mu_output_cukup_layak(z):
    return triangle(z, 20, 40, 60)

def mu_output_layak(z):
    return triangle(z, 45, 65, 85)

def mu_output_sangat_layak(z):
    return linear_ascent(z, 65, 100)

def defuzzify(alpha, n_titik=1000):
    z_min, z_max = 0.0, 100.0
    step = (z_max - z_min) / n_titik

    pembilang = 0.0
    penyebut  = 0.0

    z = z_min
    while z <= z_max:
        mu_tidak_layak = min(alpha['tidak_layak'],  mu_output_tidak_layak(z))
        mu_cukup_layak = min(alpha['cukup_layak'],  mu_output_cukup_layak(z))
        mu_layak  = min(alpha['layak'],        mu_output_layak(z))
        mu_sangat_layak = min(alpha['sangat_layak'], mu_output_sangat_layak(z))

        mu_aggregate = max(mu_tidak_layak, mu_cukup_layak, mu_layak, mu_sangat_layak)

        pembilang += z * mu_aggregate
        penyebut  += mu_aggregate
        z += step

    if penyebut == 0:
        return 0.0
    return pembilang / penyebut

def count_score(restoran):
    mu_service = fuzzify_service(restoran['pelayanan'])
    mu_price = fuzzify_price(restoran['harga'])
    alpha = inference(mu_service, mu_price)
    skor  = defuzzify(alpha)
    return round(skor, 4)

def process_all(data):
    hasil = []
    for r in data:
        score = count_score(r)
        hasil.append({
            'id'       : r['id'],
            'pelayanan': r['pelayanan'],
            'harga'    : r['harga'],
            'skor'     : score
        })
    hasil.sort(key=lambda x: x['skor'], reverse=True)
    return hasil[:5]

def save_output(top5, path_output):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Peringkat Restoran"

    headers = ["Peringkat", "ID Restoran", "Kualitas Pelayanan (1-100)",
               "Harga (Rp)", "Skor Kelayakan (0-100)"]
    worksheet.append(headers)

    for rank, r in enumerate(top5, 1):
        worksheet.append([
            rank,
            r['id'],
            r['pelayanan'],
            r['harga'],
            r['skor']
        ])
 
        worksheet.cell(row=rank + 1, column=4).number_format = '#,##0'
        worksheet.cell(row=rank + 1, column=5).number_format = '0.0000'
 
    workbook.save(path_output)
    print(f"Output disimpan ke: {path_output}")

if __name__ == "__main__":
    INPUT_FILE  = "restoran.xlsx"
    OUTPUT_FILE = "top5_restoran.xlsx"

    data = read_file(INPUT_FILE)
    top5 = process_all(data)

    print("5 Restoran Terbaik:\n")
    print(f"  {'Peringkat':<10} {'ID':>4}  {'Pelayanan':>10}  {'Harga (Rp)':>12}  {'Skor':>8}")
    print("  " + "-" * 52)
    for rank, r in enumerate(top5, 1):
        print(f"  {rank:<10} {r['id']:>4}  {r['pelayanan']:>10.0f}  {r['harga']:>12,.0f}  {r['skor']:>8.4f}")

    save_output(top5, OUTPUT_FILE)